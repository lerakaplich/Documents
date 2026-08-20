from io import BytesIO
from datetime import date, datetime, timedelta
from typing import Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from server.app.repositories.overtime_repo import OvertimeRepository
from server.app.services.common.security_service import SecurityService


class OvertimeExportService:
    def __init__(
            self,
            overtime_repo: OvertimeRepository,
            security: SecurityService
    ):
        self.overtime_repo = overtime_repo
        self.security = security

    def calculate_hours(self, start_time, end_time) -> float:
        """Вспомогательный расчет длительности в десятичных часах"""
        if not start_time or not end_time:
            return 0.0

        dt_start = datetime.combine(date.today(), start_time)
        dt_end = datetime.combine(date.today(), end_time)

        # Если переработка перешла через полночь
        if dt_end < dt_start:
            dt_end += timedelta(days=1)

        diff_seconds = (dt_end - dt_start).total_seconds()
        return round(diff_seconds / 3600.0, 2)

    def format_hours(self, decimal_hours: float) -> str:
        """Преобразование десятичных часов в формат ЧЧ:ММ"""
        total_minutes = int(round(decimal_hours * 60))
        hours = total_minutes // 60
        minutes = total_minutes % 60
        return f"{hours}:{minutes:02d}"

    async def generate_report_buffer(
            self,
            dept_id: Optional[int] = None,
            start_date: Optional[date] = None,
            end_date: Optional[date] = None
    ) -> tuple[BytesIO, str]:

        # 1. Запрашиваем записи
        rows = await self.overtime_repo.get_overtimes_for_export(
            dept_id=dept_id,
            start_date=start_date,
            end_date=end_date
        )

        # 2. Создаем Excel книгу
        wb = Workbook()
        ws = wb.active
        ws.title = "Переработки"

        # Стили
        font_title = Font(name="Calibri", size=14, bold=True)
        font_dept_header = Font(name="Calibri", size=12, bold=True, color="1B232A")
        font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        font_regular = Font(name="Calibri", size=11)
        font_bold = Font(name="Calibri", size=11, bold=True)

        header_fill = PatternFill(start_color="1B232A", end_color="1B232A", fill_type="solid")
        dept_fill = PatternFill(start_color="E6ECF2", end_color="E6ECF2", fill_type="solid")
        subtotal_fill = PatternFill(start_color="F7F9FA", end_color="F7F9FA", fill_type="solid")

        thin_border = Border(
            left=Side(style='thin', color='D3D3D3'),
            right=Side(style='thin', color='D3D3D3'),
            top=Side(style='thin', color='D3D3D3'),
            bottom=Side(style='thin', color='D3D3D3')
        )

        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # Заголовок отчета
        period_str = f"с {start_date.strftime('%d.%m.%Y')} по {end_date.strftime('%d.%m.%Y')}" if start_date and end_date else "за весь период"
        ws.append([f"Сводный отчет по переработкам ({period_str})"])
        ws.cell(row=1, column=1).font = font_title
        ws.row_dimensions[1].height = 30
        ws.append([])

        # Заголовки колонок
        headers = ["№", "ФИО", "Должность", "Дата", "Время", "Описание переработки", "Всего часов"]
        ws.append(headers)
        header_row_idx = 3
        ws.row_dimensions[header_row_idx].height = 25

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row_idx, column=col_idx)
            cell.font = font_header
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = thin_border

        # 3. Группируем данные: Dept -> Employee -> Items
        grouped_data = {}
        total_period_hours = 0.0

        for ot, last_name, first_name, patronymic, pos_name, d_id, d_name in rows:
            fio = f"{last_name or ''} {first_name or ''} {patronymic or ''}".strip()
            duration = self.calculate_hours(ot.overtime_start, ot.overtime_end)
            total_period_hours += duration

            if d_name not in grouped_data:
                grouped_data[d_name] = {}
            if fio not in grouped_data[d_name]:
                grouped_data[d_name][fio] = {"position": pos_name, "items": [], "total": 0.0}

            time_str = ""
            if ot.overtime_start and ot.overtime_end:
                time_str = f"{ot.overtime_start.strftime('%H:%M')}-{ot.overtime_end.strftime('%H:%M')}"

            grouped_data[d_name][fio]["items"].append({
                "date": ot.overtime_date.strftime('%d.%m.%Y') if ot.overtime_date else '',
                "time": time_str,
                "note": ot.note_text or "",
                "duration": duration
            })
            grouped_data[d_name][fio]["total"] += duration

        # 4. Вывод данных с объединением ячеек
        for d_name, employees in grouped_data.items():
            # Шапка подразделения
            ws.append([f"Подразделение: {d_name}"])
            dept_row = ws.max_row
            ws.merge_cells(start_row=dept_row, start_column=1, end_row=dept_row, end_column=7)
            ws.cell(row=dept_row, column=1).font = font_dept_header
            ws.cell(row=dept_row, column=1).fill = dept_fill
            ws.row_dimensions[dept_row].height = 24

            dept_total_hours = 0.0
            emp_idx = 1

            for fio, emp_data in employees.items():
                items = emp_data["items"]
                items_count = len(items)
                start_row = ws.max_row + 1

                for item in items:
                    ws.append([
                        emp_idx,
                        fio,
                        emp_data["position"],
                        item["date"],
                        item["time"],
                        item["note"],
                        self.format_hours(emp_data["total"])
                    ])
                    cur_row = ws.max_row
                    ws.row_dimensions[cur_row].height = 22

                    # Применяем стили ко всем ячейкам строки
                    for c in range(1, 8):
                        cell = ws.cell(row=cur_row, column=c)
                        cell.font = font_regular
                        cell.border = thin_border
                        if c in [1, 4, 5, 7]:
                            cell.alignment = align_center
                        else:
                            cell.alignment = align_left

                end_row = ws.max_row

                # Объединяем ячейки по вертикали для одного сотрудника (№, ФИО, Должность, Всего часов)
                if items_count > 1:
                    ws.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)  # №
                    ws.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=2)  # ФИО
                    ws.merge_cells(start_row=start_row, start_column=3, end_row=end_row, end_column=3)  # Должность
                    ws.merge_cells(start_row=start_row, start_column=7, end_row=end_row, end_column=7)  # Всего часов

                # Для ячейки "Всего часов" ставим жирный шрифт
                ws.cell(row=start_row, column=7).font = font_bold

                dept_total_hours += emp_data["total"]
                emp_idx += 1

            # Итого по подразделению
            ws.append(["", "Итого по подразделению:", "", "", "", "", self.format_hours(dept_total_hours)])
            d_total_row = ws.max_row
            ws.merge_cells(start_row=d_total_row, start_column=2, end_row=d_total_row, end_column=6)

            for c in range(1, 8):
                cell = ws.cell(row=d_total_row, column=c)
                cell.fill = subtotal_fill
                cell.border = thin_border
                cell.font = font_bold

            ws.cell(row=d_total_row, column=2).alignment = align_left
            ws.cell(row=d_total_row, column=7).alignment = align_center
            ws.row_dimensions[d_total_row].height = 22
            ws.append([])  # Пустая строка

        # ОБЩИЙ ИТОГ
        ws.append(["", "ОБЩИЙ ИТОГ:", "", "", "", "", self.format_hours(total_period_hours)])
        grand_row = ws.max_row
        ws.merge_cells(start_row=grand_row, start_column=2, end_row=grand_row, end_column=6)

        for c in range(1, 8):
            cell = ws.cell(row=grand_row, column=c)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=12, bold=True)

        ws.cell(row=grand_row, column=7).font = Font(name="Calibri", size=12, bold=True, color="C00000")
        ws.cell(row=grand_row, column=7).alignment = align_center
        ws.row_dimensions[grand_row].height = 25

        # Настройка ширины колонок
        col_widths = {'A': 6, 'B': 30, 'C': 26, 'D': 14, 'E': 16, 'F': 38, 'G': 16}
        for col, width in col_widths.items():
            ws.column_dimensions[col].width = width

        # Выгрузка в байты
        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)

        filename = f"overtime_report_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        return stream, filename